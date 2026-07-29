#!/usr/bin/env python3
"""
Generate school-wise Nandotsav 2026 final letters from a DOCX template and CSV data.

Outputs for each school inside ./output:
- <school_slug>.pdf
- <school_slug>_page_<n>.png (all PDF pages)
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List

import fitz  # PyMuPDF
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass
class Finalist:
    group: str
    competition: str
    name: str
    klass: str
    section: str
    rank: str


def apply_font_to_run(run, font_name: str = "Arial", font_size: int = 12) -> None:
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(font_size)


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def normalize_group(value: str) -> str:
    raw = normalize_space(value).upper().replace("-", " ")
    if not raw:
        return ""

    match = re.search(r"(?:GROUP\s*)?(\d+)", raw)
    if match:
        return f"GROUP {int(match.group(1))}"

    if raw.startswith("G") and raw[1:].isdigit():
        return f"GROUP {int(raw[1:])}"

    return raw


def group_sort_key(group_value: str) -> tuple[int, str]:
    match = re.search(r"(\d+)", group_value or "")
    if match:
        return (int(match.group(1)), group_value)
    return (999, group_value)


def rank_sort_key(rank_value: str) -> int:
    rank_value = normalize_space(rank_value)
    if rank_value.isdigit():
        return int(rank_value)
    return 999


def slugify(name: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9]+", "_", normalize_space(name))
    safe = safe.strip("_")
    return safe.lower() or "school"


def read_school_data(csv_path: Path) -> Dict[str, Dict[str, object]]:
    schools: Dict[str, Dict[str, object]] = defaultdict(lambda: {
        "display_counter": Counter(),
        "rows": [],
    })

    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            school_name = normalize_space(row.get("School", ""))
            if not school_name:
                continue

            school_key = school_name.casefold()
            schools[school_key]["display_counter"][school_name] += 1

            finalist = Finalist(
                group=normalize_group(row.get("Group", "")),
                competition=normalize_space(row.get("Comp", "")),
                name=normalize_space(row.get("Name", "")),
                klass=normalize_space(row.get("Class", "")),
                section=normalize_space(row.get("SECTION", "")),
                rank=normalize_space(row.get("Rank", "")),
            )
            schools[school_key]["rows"].append(finalist)

    resolved: Dict[str, Dict[str, object]] = {}
    for school_key, payload in schools.items():
        display_name = payload["display_counter"].most_common(1)[0][0]
        rows: List[Finalist] = payload["rows"]
        rows.sort(
            key=lambda x: (
                group_sort_key(x.group),
                x.competition.upper(),
                rank_sort_key(x.rank),
                x.name.upper(),
            )
        )
        resolved[school_key] = {
            "school_name": display_name,
            "rows": rows,
        }

    return resolved


def replace_school_placeholder(doc: Document, school_name: str) -> None:
    target = "{School Name}"
    replacement = school_name

    for para in doc.paragraphs:
        if target in para.text:
            replaced = False
            for run in para.runs:
                if target in run.text:
                    run.text = run.text.replace(target, replacement)
                    apply_font_to_run(run, "Arial", 12)
                    replaced = True
            if not replaced:
                para.text = para.text.replace(target, replacement)
                for run in para.runs:
                    if school_name in run.text:
                        apply_font_to_run(run, "Arial", 12)
            return

    raise ValueError("Could not find '{School Name}' placeholder in the template")


def insert_table_after(doc: Document, paragraph: Paragraph, rows: int, cols: int) -> Table:
    new_table = doc.add_table(rows=rows, cols=cols)
    paragraph._p.addnext(new_table._tbl)
    return new_table


def fill_annexure3_table(doc: Document, finalists: List[Finalist]) -> None:
    annexure_para = None
    for para in doc.paragraphs:
        text = normalize_space(para.text)
        if "Annexure" in text and "3" in text and "Finalist" in text:
            annexure_para = para
            break

    if annexure_para is None:
        raise ValueError("Could not find 'Annexure - 3 Finalists' heading in template")

    table = insert_table_after(doc, annexure_para, rows=1, cols=5)
    table.style = "Table Grid"

    headers = ["GROUP", "COMPETITION", "NAME", "CLASS", "SECTION"]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text

    for entry in finalists:
        row_cells = table.add_row().cells
        row_cells[0].text = entry.group
        row_cells[1].text = entry.competition
        row_cells[2].text = entry.name
        row_cells[3].text = entry.klass
        row_cells[4].text = entry.section

        # Keep finalist names consistent in output PDFs.
        for paragraph in row_cells[2].paragraphs:
            for run in paragraph.runs:
                apply_font_to_run(run, "Arial", 12)


def convert_docx_to_pdf_with_word(input_docx: Path, output_pdf: Path) -> None:
    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    applescript = f'''set inFile to POSIX file "{input_docx}"
set outFile to POSIX file "{output_pdf}"
tell application "Microsoft Word"
    activate
    open inFile
    save as active document file name outFile file format format PDF
    close active document saving no
end tell
'''
    result = subprocess.run(
        ["osascript", "-"],
        input=applescript,
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        error_text = (result.stderr or result.stdout or "Unknown AppleScript error").strip()
        raise RuntimeError(f"Word PDF export failed: {error_text}")

    if not output_pdf.exists():
        raise FileNotFoundError(f"PDF was not created: {output_pdf}")


def convert_pdf_to_images(pdf_path: Path, output_prefix: Path, scale: float = 2.0) -> int:
    doc = fitz.open(pdf_path)
    try:
        image_count = 0
        matrix = fitz.Matrix(scale, scale)
        for i, page in enumerate(doc):
            pix = page.get_pixmap(matrix=matrix)
            image_path = output_prefix.parent / f"{output_prefix.name}_page_{i + 1}.png"
            pix.save(str(image_path))
            image_count += 1
        return image_count
    finally:
        doc.close()


def generate_docx_for_school(template_path: Path, temp_dir: Path, school_name: str, finalists: List[Finalist]) -> tuple[str, Path]:
    school_slug = slugify(school_name)
    temp_dir.mkdir(parents=True, exist_ok=True)

    docx_path = temp_dir / f"{school_slug}.docx"

    doc = Document(str(template_path))
    replace_school_placeholder(doc, school_name)
    fill_annexure3_table(doc, finalists)
    doc.save(str(docx_path))
    return school_slug, docx_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate Nandotsav finalist letters per school")
    parser.add_argument(
        "--csv",
        dest="csv_path",
        type=Path,
        default=Path("Nandotsav 2026 - Top3.csv"),
        help="Path to Top3 CSV file",
    )
    parser.add_argument(
        "--template",
        dest="template_path",
        type=Path,
        default=Path("Nandotsav 2026 Finals.docx"),
        help="Path to DOCX template",
    )
    parser.add_argument(
        "--output",
        dest="output_dir",
        type=Path,
        default=Path("output"),
        help="Output directory",
    )
    parser.add_argument(
        "--clean",
        action="store_true",
        help="Delete output directory before generating files",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    csv_path = args.csv_path.resolve()
    template_path = args.template_path.resolve()
    output_dir = args.output_dir.resolve()

    if not csv_path.exists():
        print(f"CSV not found: {csv_path}", file=sys.stderr)
        return 1

    if not template_path.exists():
        print(f"Template not found: {template_path}", file=sys.stderr)
        return 1

    if args.clean and output_dir.exists():
        shutil.rmtree(output_dir)

    school_data = read_school_data(csv_path)
    if not school_data:
        print("No school data found in CSV.", file=sys.stderr)
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)
    temp_dir = Path(tempfile.mkdtemp(prefix="tmp_docx_", dir=output_dir))

    failures = []
    generated: List[tuple[str, str, Path, Path]] = []

    for payload in school_data.values():
        school_name = payload["school_name"]
        finalists = payload["rows"]
        try:
            school_slug, docx_path = generate_docx_for_school(
                template_path,
                temp_dir,
                school_name,
                finalists,
            )
            pdf_path = output_dir / f"{school_slug}.pdf"
            generated.append((school_name, school_slug, docx_path, pdf_path))
        except Exception as exc:  # noqa: BLE001
            failures.append((school_name, str(exc)))
            print(f"Failed: {school_name} -> {exc}", file=sys.stderr)

    conversion_success = 0
    for school_name, school_slug, docx_path, pdf_path in generated:
        try:
            convert_docx_to_pdf_with_word(docx_path, pdf_path)
            pages = convert_pdf_to_images(pdf_path, output_dir / school_slug)
            conversion_success += 1

            print(f"Generated: {school_name}")
            print(f"  PDF : {pdf_path}")
            print(f"  PNG : {pages} page image(s)")
        except Exception as exc:  # noqa: BLE001
            failures.append((school_name, str(exc)))
            print(f"Failed: {school_name} -> {exc}", file=sys.stderr)

    shutil.rmtree(temp_dir, ignore_errors=True)

    print("\nSummary")
    print(f"  Schools processed : {len(school_data)}")
    print(f"  Success           : {conversion_success}")
    print(f"  Failed            : {len(failures)}")

    if failures:
        print("\nFailures:", file=sys.stderr)
        for school_name, reason in failures:
            print(f"  - {school_name}: {reason}", file=sys.stderr)
        return 2

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
