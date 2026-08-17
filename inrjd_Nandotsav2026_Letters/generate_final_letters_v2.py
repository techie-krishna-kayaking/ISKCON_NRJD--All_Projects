#!/usr/bin/env python3
"""
Generate school-wise Nandotsav 2026 final letters from a DOCX template and Excel data.

Outputs for each school inside workflow subfolders:
- ./output/finalist/...
- ./output/medalist/...
- <school_slug>.pdf
- <school_slug>_page_<n>.png (all PDF pages)

Supports: macOS (Word or LibreOffice), Windows (Word or LibreOffice), Linux (LibreOffice)
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

import fitz  # PyMuPDF
import yaml
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook


@dataclass
class Finalist:
    group: str
    competition: str
    name: str
    klass: str
    section: str
    rank: str


@dataclass
class FancyParticipant:
    name: str
    klass: str
    section: str


@dataclass
class MedalistWinner:
    group: str
    competition: str
    name: str
    klass: str
    section: str


def apply_font_to_run(run, font_name: str = "Arial", font_size: int = 12) -> None:
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(font_size)


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def normalize_group(value: str) -> str:
    raw = normalize_space(value).upper().replace("-", " ")
    if not raw:
        return ""

    match = re.search(r"(?:GROUP\s*)?(\d+)", raw)
    if match:
        return f"GROUP {int(match.group(1))}"

    if raw.startswith("G") and raw[1:].isdigit():
        return f"GROUP {int(raw[1:])}"

    return raw


def group_sort_key(group_value: str) -> tuple[int, str]:
    match = re.search(r"(\d+)", group_value or "")
    if match:
        return (int(match.group(1)), group_value)
    return (999, group_value)


def rank_sort_key(rank_value: str) -> int:
    rank_value = normalize_space(rank_value)
    if rank_value.isdigit():
        return int(rank_value)
    return 999


def slugify(name: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9]+", "_", normalize_space(name))
    safe = safe.strip("_")
    return safe.lower() or "school"


def normalize_header(value: str) -> str:
    text = normalize_space(value)
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def pick_row_value(row: Dict[str, str], *aliases: str) -> str:
    for alias in aliases:
        key = normalize_header(alias)
        value = row.get(key, "")
        if normalize_space(value):
            return normalize_space(value)
    return ""


def read_sheet_rows(excel_path: Path, sheet_name: str) -> List[Dict[str, str]]:
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        target_sheet = None
        for actual_name in wb.sheetnames:
            if actual_name.casefold() == sheet_name.casefold():
                target_sheet = actual_name
                break

        if target_sheet is None:
            available = ", ".join(wb.sheetnames)
            raise ValueError(
                f"Sheet '{sheet_name}' not found in {excel_path.name}. "
                f"Available sheets: {available}"
            )

        ws = wb[target_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return []

        headers = [normalize_header(str(col or "")) for col in header_row]
        records: List[Dict[str, str]] = []
        for row in rows_iter:
            record: Dict[str, str] = {}
            non_empty = False
            for idx, cell in enumerate(row):
                if idx >= len(headers) or not headers[idx]:
                    continue
                value = normalize_space(str(cell or ""))
                if value:
                    non_empty = True
                record[headers[idx]] = value
            if non_empty:
                records.append(record)
        return records
    finally:
        wb.close()


def read_school_data_from_excel(
    excel_path: Path,
    finalists_sheet: str,
    fancydress_sheet: str,
) -> Dict[str, Dict[str, object]]:
    schools: Dict[str, Dict[str, object]] = defaultdict(lambda: {
        "display_counter": Counter(),
        "finalists": [],
        "fancy_participants": [],
    })

    finalist_rows = read_sheet_rows(excel_path, finalists_sheet)
    fancy_rows = read_sheet_rows(excel_path, fancydress_sheet)

    for row in finalist_rows:
        school_name = pick_row_value(row, "SchoolName", "School Name", "School")
        if not school_name:
            continue

        school_key = school_name.casefold()
        schools[school_key]["display_counter"][school_name] += 1

        finalist = Finalist(
            group=normalize_group(pick_row_value(row, "Group")),
            competition=pick_row_value(row, "Competition"),
            name=pick_row_value(row, "StudentName", "Student Name", "Name"),
            klass=pick_row_value(row, "Class", "Klass"),
            section=pick_row_value(row, "Sec", "Section"),
            rank=pick_row_value(row, "Rank"),
        )
        schools[school_key]["finalists"].append(finalist)

    for row in fancy_rows:
        school_name = pick_row_value(row, "SchoolName", "School Name", "School")
        if not school_name:
            continue

        school_key = school_name.casefold()
        schools[school_key]["display_counter"][school_name] += 1

        participant = FancyParticipant(
            name=pick_row_value(row, "StudentName", "Student Name", "Participant", "Name"),
            klass=pick_row_value(row, "Class", "Klass"),
            section=pick_row_value(row, "Sec", "Section"),
        )
        if participant.name:
            schools[school_key]["fancy_participants"].append(participant)

    resolved: Dict[str, Dict[str, object]] = {}
    for school_key, payload in schools.items():
        display_name = payload["display_counter"].most_common(1)[0][0]
        finalists: List[Finalist] = payload["finalists"]
        finalists.sort(
            key=lambda x: (
                group_sort_key(x.group),
                x.competition.upper(),
                rank_sort_key(x.rank),
                x.name.upper(),
            )
        )
        fancy_participants: List[FancyParticipant] = payload["fancy_participants"]
        fancy_participants.sort(key=lambda x: (x.name.upper(), x.klass.upper(), x.section.upper()))

        resolved[school_key] = {
            "school_name": display_name,
            "finalists": finalists,
            "fancy_participants": fancy_participants,
        }

    return resolved


def replace_school_placeholder(doc: Document, school_name: str, font_size: int = 12) -> None:
    target = "{School Name}"
    replacement = school_name

    for para in doc.paragraphs:
        if target in para.text:
            replaced = False
            for run in para.runs:
                if target in run.text:
                    run.text = run.text.replace(target, replacement)
                    apply_font_to_run(run, "Arial", font_size)
                    replaced = True
            if not replaced:
                para.text = para.text.replace(target, replacement)
                for run in para.runs:
                    if school_name in run.text:
                        apply_font_to_run(run, "Arial", font_size)
            return

    raise ValueError("Could not find '{School Name}' placeholder in the template")


def insert_table_after(doc: Document, paragraph: Paragraph, rows: int, cols: int) -> Table:
    new_table = doc.add_table(rows=rows, cols=cols)
    paragraph._p.addnext(new_table._tbl)
    return new_table


def find_annexure_heading(doc: Document, include_terms: List[str]) -> Optional[Paragraph]:
    found = None
    for para in doc.paragraphs:
        text = normalize_space(para.text).lower()
        if all(term in text for term in include_terms):
            found = para
    return found


def fill_annexure3_table(doc: Document, finalists: List[Finalist]) -> None:
    annexure_para = find_annexure_heading(doc, ["annexure", "3", "final"])
    if annexure_para is None:
        raise ValueError("Could not find 'ANNEXURE - 3 FINALISTS' heading in template")

    table = insert_table_after(doc, annexure_para, rows=1, cols=5)
    table.style = "Table Grid"

    headers = ["GROUP", "COMPETITION", "NAME", "CLASS", "SECTION"]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text

    for entry in finalists:
        row_cells = table.add_row().cells
        row_cells[0].text = entry.group
        row_cells[1].text = entry.competition
        row_cells[2].text = entry.name
        row_cells[3].text = entry.klass
        row_cells[4].text = entry.section

        # Keep finalist names consistent in output PDFs.
        for paragraph in row_cells[2].paragraphs:
            for run in paragraph.runs:
                apply_font_to_run(run, "Arial", 12)


def fill_annexure4_table(doc: Document, participants: List[FancyParticipant]) -> None:
    annexure_para = find_annexure_heading(doc, ["annexure", "4", "fancy", "dress"])
    if annexure_para is None:
        raise ValueError("Could not find 'ANNEXURE - 4 FANCY DRESS COMPETITION PARTICIPANTS' heading in template")

    table = insert_table_after(doc, annexure_para, rows=1, cols=3)
    table.style = "Table Grid"

    headers = ["NAME", "CLASS", "SECTION"]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text

    for entry in participants:
        row_cells = table.add_row().cells
        row_cells[0].text = entry.name
        row_cells[1].text = entry.klass
        row_cells[2].text = entry.section

        for paragraph in row_cells[0].paragraphs:
            for run in paragraph.runs:
                apply_font_to_run(run, "Arial", 12)


def read_medalist_data_from_excel(
    excel_path: Path,
    overall_sheet: str,
) -> Dict[str, Dict[str, object]]:
    schools: Dict[str, Dict[str, object]] = defaultdict(lambda: {
        "display_counter": Counter(),
        "overall_winners": [],
    })

    overall_rows = read_sheet_rows(excel_path, overall_sheet)

    for row in overall_rows:
        school_name = pick_row_value(row, "Schoolname", "SchoolName", "School Name", "School")
        if not school_name:
            continue

        winner = MedalistWinner(
            group=normalize_group(pick_row_value(row, "Group")),
            competition=pick_row_value(row, "Competition"),
            name=pick_row_value(row, "Student Name", "StudentName", "Name"),
            klass=pick_row_value(row, "Class", "Klass"),
            section=pick_row_value(row, "Sec", "Section"),
        )
        if not winner.name:
            continue

        school_key = school_name.casefold()
        schools[school_key]["display_counter"][school_name] += 1
        schools[school_key]["overall_winners"].append(winner)

    resolved: Dict[str, Dict[str, object]] = {}
    for school_key, payload in schools.items():
        display_name = payload["display_counter"].most_common(1)[0][0]
        overall_winners: List[MedalistWinner] = payload["overall_winners"]

        overall_winners.sort(
            key=lambda x: (
                group_sort_key(x.group),
                x.competition.upper(),
                x.name.upper(),
                x.klass.upper(),
                x.section.upper(),
            )
        )

        resolved[school_key] = {
            "school_name": display_name,
            "overall_winners": overall_winners,
        }

    return resolved


def add_medalist_table(doc: Document, winners: List[MedalistWinner]) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["GROUP", "COMPETITION", "STUDENT NAME", "CLASS", "SEC"]

    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                apply_font_to_run(run, "Arial", 11)

    for winner in winners:
        row_cells = table.add_row().cells
        row_cells[0].text = winner.group
        row_cells[1].text = winner.competition
        row_cells[2].text = winner.name
        row_cells[3].text = winner.klass
        row_cells[4].text = winner.section
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    apply_font_to_run(run, "Arial", 11)


def add_medalist_section(
    doc: Document,
    heading_text: str,
    winners: List[MedalistWinner],
) -> None:
    doc.add_page_break()

    heading_para = doc.add_paragraph()
    heading_run = heading_para.add_run(heading_text)
    heading_run.bold = True
    apply_font_to_run(heading_run, "Arial", 14)

    add_medalist_table(doc, winners)


def generate_medalist_docx_for_school(
    template_path: Path,
    temp_dir: Path,
    school_name: str,
    overall_winners: List[MedalistWinner],
) -> tuple[str, Path]:
    school_slug = slugify(school_name)
    temp_dir.mkdir(parents=True, exist_ok=True)

    docx_path = temp_dir / f"{school_slug}.docx"
    doc = Document(str(template_path))
    replace_school_placeholder(doc, school_name, font_size=11)
    add_medalist_section(doc, "Inter-School Winners - ", overall_winners)

    doc.save(str(docx_path))
    return school_slug, docx_path


def convert_docx_to_pdf_with_word_mac(input_docx: Path, output_pdf: Path) -> None:
    """Convert DOCX to PDF using Microsoft Word on macOS via AppleScript."""
    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    applescript = f'''set inFile to POSIX file "{input_docx}"
set outFile to POSIX file "{output_pdf}"
tell application "Microsoft Word"
    activate
    open inFile
    save as active document file name outFile file format format PDF
    close active document saving no
end tell
'''
    result = subprocess.run(
        ["osascript", "-"],
        input=applescript,
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        error_text = (result.stderr or result.stdout or "Unknown AppleScript error").strip()
        raise RuntimeError(f"Word PDF export failed: {error_text}")

    if not output_pdf.exists():
        raise FileNotFoundError(f"PDF was not created: {output_pdf}")


def convert_docx_to_pdf_with_word_windows(input_docx: Path, output_pdf: Path) -> None:
    """Convert DOCX to PDF using Microsoft Word on Windows via COM."""
    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    try:
        import win32com.client  # type: ignore
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        try:
            doc = word.Documents.Open(str(input_docx.resolve()))
            doc.SaveAs(str(output_pdf.resolve()), FileFormat=17)  # 17 = wdFormatPDF
            doc.Close()
        finally:
            word.Quit()
    except ImportError:
        raise RuntimeError(
            "pywin32 not installed. Install with: pip install pywin32\n"
            "Or install LibreOffice and use --use-libreoffice flag"
        )
    except Exception as e:
        raise RuntimeError(f"Word PDF export failed: {e}")

    if not output_pdf.exists():
        raise FileNotFoundError(f"PDF was not created: {output_pdf}")


def convert_docx_to_pdf_with_libreoffice(input_docx: Path, output_pdf: Path) -> None:
    """Convert DOCX to PDF using LibreOffice (cross-platform)."""
    output_pdf.parent.mkdir(parents=True, exist_ok=True)

    # Find soffice executable - comprehensive path list for all platforms
    soffice_paths = [
        # macOS
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        # Windows (standard 64-bit)
        Path("C:\\Program Files\\LibreOffice\\program\\soffice.exe"),
        # Windows (32-bit)
        Path("C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe"),
        # Windows (user/portable)
        Path(Path.home() / "AppData" / "Local" / "Programs" / "LibreOffice" / "program" / "soffice.exe"),
        # Linux
        Path("/usr/bin/soffice"),
        Path("/usr/bin/libreoffice"),
        # Linux Snap
        Path("/snap/bin/libreoffice"),
    ]

    soffice = None
    for path in soffice_paths:
        if path.exists():
            soffice = path
            break

    if soffice is None:
        if sys.platform == "darwin":
            raise RuntimeError(
                "LibreOffice not found. Install with: brew install libreoffice\n"
                "Or download from: https://www.libreoffice.org"
            )
        elif sys.platform == "win32":
            raise RuntimeError(
                "LibreOffice not found. Download from: https://www.libreoffice.org\n"
                "Standard installation path: C:\\Program Files\\LibreOffice"
            )
        else:
            raise RuntimeError(
                "LibreOffice not found. Install with your package manager:\n"
                "  • Ubuntu/Debian: sudo apt install libreoffice\n"
                "  • Fedora: sudo dnf install libreoffice\n"
                "  • Arch: sudo pacman -S libreoffice-still\n"
                "  • Snap: sudo snap install libreoffice"
            )

    try:
        subprocess.run(
            [
                str(soffice),
                "--headless",
                "--convert-to", "pdf",
                "--outdir", str(output_pdf.parent),
                str(input_docx.resolve()),
            ],
            capture_output=True,
            text=True,
            check=True,
            timeout=60,
        )
    except subprocess.CalledProcessError as e:
        raise RuntimeError(f"LibreOffice PDF conversion failed: {e.stderr or e.stdout}")
    except subprocess.TimeoutExpired:
        raise RuntimeError("LibreOffice PDF conversion timed out")

    # LibreOffice saves as input_filename.pdf in the output directory
    expected_pdf = output_pdf.parent / f"{input_docx.stem}.pdf"
    if expected_pdf.exists() and expected_pdf != output_pdf:
        expected_pdf.rename(output_pdf)
    elif not output_pdf.exists():
        raise FileNotFoundError(f"PDF was not created: {output_pdf}")


def convert_pdf_to_images(pdf_path: Path, output_prefix: Path, scale: float = 2.0) -> int:
    doc = fitz.open(pdf_path)
    try:
        image_count = 0
        matrix = fitz.Matrix(scale, scale)
        for i, page in enumerate(doc):
            pix = page.get_pixmap(matrix=matrix)
            image_path = output_prefix.parent / f"{output_prefix.name}_page_{i + 1}.png"
            pix.save(str(image_path))
            image_count += 1
        return image_count
    finally:
        doc.close()


def generate_docx_for_school(
    template_path: Path,
    temp_dir: Path,
    school_name: str,
    finalists: List[Finalist],
    fancy_participants: List[FancyParticipant],
) -> tuple[str, Path]:
    school_slug = slugify(school_name)
    temp_dir.mkdir(parents=True, exist_ok=True)

    docx_path = temp_dir / f"{school_slug}.docx"

    doc = Document(str(template_path))
    replace_school_placeholder(doc, school_name)
    fill_annexure3_table(doc, finalists)
    fill_annexure4_table(doc, fancy_participants)
    doc.save(str(docx_path))
    return school_slug, docx_path


def detect_pdf_converter() -> Optional[str]:
    """Detect available PDF converter: 'word', 'libreoffice', or None."""
    if sys.platform == "darwin":  # macOS
        # Check for Microsoft Word on macOS
        word_app_path = Path("/Applications/Microsoft Word.app")
        if word_app_path.exists():
            return "word"
    
    if sys.platform == "win32":  # Windows
        # Check for Microsoft Word on Windows
        try:
            import win32com.client  # type: ignore
            word = win32com.client.Dispatch("Word.Application")
            word.Quit()
            return "word"
        except (ImportError, Exception):
            pass

    # Check for LibreOffice on all platforms
    soffice_paths = [
        # macOS
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        # Windows (standard installations)
        Path("C:\\Program Files\\LibreOffice\\program\\soffice.exe"),
        Path("C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe"),
        # Windows (user installations)
        Path(Path.home() / "AppData" / "Local" / "Programs" / "LibreOffice" / "program" / "soffice.exe"),
        # Linux
        Path("/usr/bin/soffice"),
        Path("/usr/bin/libreoffice"),
        Path("/snap/bin/libreoffice"),  # Snap installation on Linux
    ]
    for path in soffice_paths:
        if path.exists():
            return "libreoffice"

    return None


def load_config(config_path: str = "config.yaml") -> Dict:
    """Load configuration from YAML file."""
    config_file = Path(config_path)
    
    # Default configuration
    default_config = {
        "workflow": "finalist",
        "finalist_excel_file": "TMPLT_FInalist_Nandotsav2026.xlsx",
        "finalist_template_file": "TMPLT_Finalist_Nandotsav2026.docx",
        "finalists_sheet": "top3",
        "fancydress_sheet": "fancydress",
        "medalist_excel_file": "TMPLT_Certificate_Nandotsav2026.xlsx",
        "medalist_template_file": "TMPLT_Medalist_Nandotsav2026.docx",
        "overall_sheet": "Overall",
        "output_folder": "output",
        "pdf_converter": "word",
        "verbose": False,
        "clean": False,
    }
    
    if not config_file.exists():
        # Create default config if it doesn't exist
        with open(config_file, "w") as f:
            yaml.dump(default_config, f, default_flow_style=False)
        return default_config
    
    try:
        with open(config_file, "r") as f:
            loaded_config = yaml.safe_load(f)
            if loaded_config:
                default_config.update(loaded_config)
    except Exception as e:
        print(f"⚠️  Warning: Could not load config.yaml: {e}", file=sys.stderr)
        print(f"   Using default configuration", file=sys.stderr)
    
    return default_config


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate Nandotsav finalist letters per school")
    parser.add_argument(
        "--workflow",
        dest="workflow",
        choices=["finalist", "medalist"],
        help="Workflow to run: finalist or medalist (default from config.yaml)",
    )
    parser.add_argument(
        "--config",
        dest="config_file",
        type=str,
        default="config.yaml",
        help="Path to YAML config file (default: config.yaml)",
    )
    parser.add_argument(
        "--excel",
        dest="excel_file",
        type=str,
        help="Excel file name inside input/ (overrides config.yaml)",
    )
    parser.add_argument(
        "--csv",
        dest="csv_file",
        type=str,
        help="Deprecated alias for --excel (kept for backward compatibility)",
    )
    parser.add_argument(
        "--finalists-sheet",
        dest="finalists_sheet",
        type=str,
        help="Sheet name for finalists data (default from config.yaml)",
    )
    parser.add_argument(
        "--fancydress-sheet",
        dest="fancydress_sheet",
        type=str,
        help="Sheet name for fancy dress participants (default from config.yaml)",
    )
    parser.add_argument(
        "--overall-sheet",
        dest="overall_sheet",
        type=str,
        help="Sheet name for inter-school winners (medalist workflow)",
    )
    parser.add_argument(
        "--template",
        dest="template_file",
        type=str,
        help="Template DOCX file name (overrides config.yaml)",
    )
    parser.add_argument(
        "--output",
        dest="output_folder",
        type=str,
        help="Output folder name (overrides config.yaml)",
    )
    parser.add_argument(
        "--clean",
        action="store_true",
        help="Delete output folder before generating files",
    )
    parser.add_argument(
        "--use-libreoffice",
        action="store_true",
        help="Force use of LibreOffice instead of Word",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Print detailed progress information",
    )
    return parser.parse_args()



def main() -> int:
    args = parse_args()

    # Load configuration from YAML
    config = load_config(args.config_file)

    # Override config with command-line arguments
    if args.workflow:
        config["workflow"] = args.workflow
    if args.finalists_sheet:
        config["finalists_sheet"] = args.finalists_sheet
    if args.fancydress_sheet:
        config["fancydress_sheet"] = args.fancydress_sheet
    if args.overall_sheet:
        config["overall_sheet"] = args.overall_sheet
    if args.output_folder:
        config["output_folder"] = args.output_folder
    if args.verbose:
        config["verbose"] = True
    if args.clean:
        config["clean"] = True
    if args.use_libreoffice:
        config["pdf_converter"] = "libreoffice"

    workflow = normalize_space(str(config.get("workflow", "finalist"))).casefold()
    if workflow not in {"finalist", "medalist"}:
        print("❌ Invalid workflow. Allowed values: finalist, medalist", file=sys.stderr)
        return 1

    if workflow == "finalist":
        excel_file_name = str(config.get("finalist_excel_file") or config.get("excel_file") or "")
        template_file_name = str(config.get("finalist_template_file") or config.get("template_file") or "")
    else:
        excel_file_name = str(config.get("medalist_excel_file") or config.get("excel_file") or "")
        template_file_name = str(config.get("medalist_template_file") or config.get("template_file") or "")

    if args.excel_file:
        excel_file_name = args.excel_file
    if args.csv_file:
        excel_file_name = args.csv_file
    if args.template_file:
        template_file_name = args.template_file

    # Set up paths - input files are in 'input/' folder
    input_dir = Path("input")
    excel_path = (input_dir / excel_file_name).resolve()
    template_path = (input_dir / template_file_name).resolve()
    output_root_dir = Path(config["output_folder"]).resolve()
    output_dir = output_root_dir / workflow

    if not input_dir.exists():
        print(f"❌ Input directory not found: {input_dir}", file=sys.stderr)
        print(f"   Please create an 'input' folder and add your Excel and DOCX files there", file=sys.stderr)
        return 1

    if not excel_path.exists():
        print(f"❌ Excel file not found: {excel_path}", file=sys.stderr)
        return 1

    if not template_path.exists():
        print(f"❌ Template file not found: {template_path}", file=sys.stderr)
        return 1

    # Determine PDF converter
    if config["pdf_converter"] == "libreoffice":
        converter = "libreoffice"
    else:
        converter = detect_pdf_converter()

    if converter is None:
        if sys.platform == "darwin":  # macOS
            msg = (
                "❌ No PDF converter found!\n\n"
                "Please install one of the following:\n"
                "  • Microsoft Word: Available from Microsoft 365 or Mac App Store\n"
                "  • LibreOffice: brew install libreoffice\n\n"
                "Or use the fallback:\n"
                "  python3 generate_final_letters_v2.py --use-libreoffice"
            )
        elif sys.platform == "win32":  # Windows
            msg = (
                "❌ No PDF converter found!\n\n"
                "Please install one of the following:\n"
                "  • Microsoft Word: Download from https://www.microsoft.com/microsoft-365\n"
                "  • LibreOffice: Download from https://www.libreoffice.org\n\n"
                "Or use the fallback:\n"
                "  python generate_final_letters_v2.py --use-libreoffice"
            )
        else:  # Linux
            msg = (
                "❌ No PDF converter found!\n\n"
                "Please install LibreOffice:\n"
                "  • Ubuntu/Debian: sudo apt install libreoffice\n"
                "  • Fedora: sudo dnf install libreoffice\n"
                "  • Arch: sudo pacman -S libreoffice-still\n"
                "  • Snap: sudo snap install libreoffice"
            )
        print(msg, file=sys.stderr)
        return 1

    if config["verbose"]:
        print(f"📋 Configuration loaded from: {args.config_file}")
        print(f"   Workflow: {workflow}")
        print(f"   Excel: {excel_path}")
        if workflow == "finalist":
            print(f"   Finalists sheet: {config['finalists_sheet']}")
            print(f"   Fancydress sheet: {config['fancydress_sheet']}")
        else:
            print(f"   Overall sheet: {config['overall_sheet']}")
        print(f"   Template: {template_path}")
        print(f"   Output root: {output_root_dir}")
        print(f"   Output: {output_dir}")
        print(f"   Converter: {converter}")
        print()

    if config["clean"] and output_dir.exists():
        shutil.rmtree(output_dir)

    if workflow == "finalist":
        school_data = read_school_data_from_excel(
            excel_path,
            config["finalists_sheet"],
            config["fancydress_sheet"],
        )
    else:
        school_data = read_medalist_data_from_excel(
            excel_path,
            config["overall_sheet"],
        )
    if not school_data:
        print("❌ No school data found in the configured Excel sheets.", file=sys.stderr)
        return 1

    if config["verbose"]:
        print(f"📊 Found {len(school_data)} schools")

    output_dir.mkdir(parents=True, exist_ok=True)
    temp_dir = Path(tempfile.mkdtemp(prefix="tmp_docx_", dir=output_dir))

    failures = []
    generated: List[tuple[str, str, Path, Path]] = []

    for payload in school_data.values():
        school_name = payload["school_name"]
        try:
            if workflow == "finalist":
                finalists = payload["finalists"]
                fancy_participants = payload["fancy_participants"]
                school_slug, docx_path = generate_docx_for_school(
                    template_path,
                    temp_dir,
                    school_name,
                    finalists,
                    fancy_participants,
                )
            else:
                overall_winners = payload["overall_winners"]
                school_slug, docx_path = generate_medalist_docx_for_school(
                    template_path,
                    temp_dir,
                    school_name,
                    overall_winners,
                )
            pdf_path = output_dir / f"{school_slug}.pdf"
            generated.append((school_name, school_slug, docx_path, pdf_path))
        except Exception as exc:  # noqa: BLE001
            failures.append((school_name, str(exc)))
            print(f"❌ Failed: {school_name} -> {exc}", file=sys.stderr)

    conversion_success = 0
    for school_name, school_slug, docx_path, pdf_path in generated:
        try:
            if converter == "word" and sys.platform == "win32":
                convert_docx_to_pdf_with_word_windows(docx_path, pdf_path)
            elif converter == "word" and sys.platform == "darwin":
                convert_docx_to_pdf_with_word_mac(docx_path, pdf_path)
            else:
                convert_docx_to_pdf_with_libreoffice(docx_path, pdf_path)

            pages = convert_pdf_to_images(pdf_path, output_dir / school_slug)
            conversion_success += 1

            print(f"✅ Generated: {school_name}")
            if config["verbose"]:
                print(f"   PDF: {pdf_path}")
                print(f"   PNG: {pages} page image(s)")
        except Exception as exc:  # noqa: BLE001
            failures.append((school_name, str(exc)))
            print(f"❌ Failed: {school_name} -> {exc}", file=sys.stderr)

    shutil.rmtree(temp_dir, ignore_errors=True)

    print("\n" + "=" * 50)
    print("Summary")
    print("=" * 50)
    print(f"Schools processed : {len(school_data)}")
    print(f"Success           : {conversion_success}")
    print(f"Failed            : {len(failures)}")

    if failures:
        print("\nFailures:", file=sys.stderr)
        for school_name, reason in failures:
            print(f"  - {school_name}: {reason}", file=sys.stderr)
        return 2

    print("\n✅ All letters generated successfully!")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
